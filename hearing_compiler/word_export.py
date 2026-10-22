"""Word export of the same reviewed content used for copy and Markdown export.

The DOCX is a readable copy, not a workspace backup: source audio and the
editable review history remain in the saved local workspace.
"""

from __future__ import annotations

import io
import re

from .pipeline import InputError, markdown

_INVALID_XML_TEXT = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _text(line: str) -> str:
    """Undo only the escapes produced by the canonical Markdown exporter."""
    return re.sub(r"\\([\\`*_{}\[\]<>#|])", r"\1", line)


def docx_bytes(draft: dict) -> bytes:
    """Return an editable Word document without changing the saved draft.

    Import the export dependency on demand so an older source installation
    can still launch, review, and copy notes before installing it.
    """
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise InputError(
            "Word export needs python-docx. Install the project's requirements.txt "
            "with your Python environment, then try exporting again. Your saved "
            "notes are unchanged; copy and Markdown export remain available."
        ) from exc

    note = markdown(draft)
    if _INVALID_XML_TEXT.search(note):
        raise InputError(
            "The note contains a control character that Word cannot store. Your "
            "saved notes are unchanged; remove that character or use Markdown export."
        )
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    styles = document.styles
    for name, size in (("Normal", 11), ("Title", 18), ("Heading 1", 13),
                       ("Heading 2", 12), ("Heading 3", 11), ("List Bullet", 11)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.widow_control = True
    normal = styles["Normal"].paragraph_format
    normal.space_after = Pt(6)
    normal.line_spacing = 1.08
    styles["Title"].paragraph_format.space_after = Pt(10)
    title_properties = styles["Title"]._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[name].paragraph_format.keep_with_next = True
        styles[name].paragraph_format.space_before = Pt(12)
        styles[name].paragraph_format.space_after = Pt(6)
    styles["List Bullet"].paragraph_format.space_after = Pt(4)

    properties = document.core_properties
    properties.title = draft["hearing"]["title"][:255]  # Word's property limit; the heading keeps the full title.
    properties.subject = "Hearing notes with review state and transcript evidence"
    properties.author = properties.last_modified_by = ""
    properties.comments = ""

    # The canonical exporter already includes review warnings, corrected
    # values, generic speaker labels, source references, and all associated
    # and unassigned Q&A. Formatting must never create a second note schema.
    for line in note.splitlines():
        if not line or line == "---":
            continue
        heading = re.match(r"^(#{1,4}) (.*)$", line)
        if heading:
            level = len(heading.group(1))
            style = "Title" if level == 1 else f"Heading {level - 1}"
            document.add_paragraph(_text(heading.group(2)), style)
        elif line.startswith("- "):
            document.add_paragraph(_text(line[2:]), "List Bullet")
        elif line.startswith("**") and line.endswith("**"):
            paragraph = document.add_paragraph()
            paragraph.add_run(_text(line[2:-2])).bold = True
        else:
            paragraph = document.add_paragraph(_text(line))
            if line == "Main Arguments:":
                paragraph.paragraph_format.keep_with_next = True

    target = io.BytesIO()
    try:
        document.save(target)
    except (TypeError, ValueError) as exc:
        raise InputError(
            "The note contains text that Word cannot store. Your saved notes are "
            "unchanged; review the note text or use Markdown export."
        ) from exc
    return target.getvalue()
