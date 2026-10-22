from __future__ import annotations

import copy
import io
import json
import re
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn

from hearing_compiler.pipeline import InputError, correct, extract, fields, markdown
from hearing_compiler.word_export import docx_bytes

ROOT = Path(__file__).resolve().parent.parent


def sample_draft():
    source = json.loads((ROOT / "fixtures/transcripts/sample-hearing.json").read_text(encoding="utf-8"))
    return extract(source)


def content(payload):
    return "\n".join(p.text for p in Document(io.BytesIO(payload)).paragraphs)


class WordExportTests(unittest.TestCase):
    def test_keeps_draft_warning_source_times_speakers_questions_and_answers(self):
        draft = sample_draft()
        original = copy.deepcopy(draft)
        text = content(docx_bytes(draft))
        for expected in (
            "Synthetic test hearing", "Review: Unresolved", "Tags: Hearing Notes; Example Committee",
            "Source reference:", "Main Arguments:", "Legislator Questions", "Question:", "Response:",
            "Unknown [unresolved]", "[draft; review needed]", "01:24:12", "Speaker 3", "s03",
            "without replacing existing work",
        ):
            self.assertIn(expected, text)
        self.assertEqual(draft, original)

    def test_every_canonical_note_line_survives_formatting(self):
        # Multiple bills, unresolved attribution and review issues must not be
        # accidentally dropped when the shorter presentation is formatted.
        source = json.loads((ROOT / "fixtures/transcripts/mixed-position.json").read_text(encoding="utf-8"))
        draft = extract(source)
        text = content(docx_bytes(draft))
        self.assertIn("Questions needing testimony or bill attribution", text)
        for line in markdown(draft).splitlines():
            if not line or line == "---":
                continue
            line = re.sub(r"^(?:#{1,4} |\- )", "", line)
            if line.startswith("**") and line.endswith("**"):
                line = line[2:-2]
            line = re.sub(r"\\([\\`*_{}\[\]<>#|])", r"\1", line)
            self.assertIn(line, text)

    def test_approved_corrected_unicode_and_markdown_characters_remain_literal(self):
        draft = sample_draft()
        changes = [{"path": path, "review": "APPROVED"} for path, item in fields(draft) if item["value"]]
        draft = correct(draft, changes + [{"path": "/questions/0/asker", "value": "Rep. José [Avery] — _reviewed_", "review": "APPROVED"}])
        text = content(docx_bytes(draft))
        self.assertIn("Review: Approved", text)
        self.assertIn("Rep. José [Avery] — _reviewed_ [human corrected]", text)
        self.assertNotIn("[draft; review needed]", text)
        self.assertNotIn("\\[Avery\\]", text)

    def test_semantic_styles_readable_page_size_and_long_content(self):
        draft = sample_draft()
        long_argument = "A family may need clear express authority for planning, while the court retains its role. " * 60
        draft = correct(draft, [{"path": "/testimonies/0/statements/0/arguments/0", "value": long_argument}])
        document = Document(io.BytesIO(docx_bytes(draft)))
        self.assertEqual(document.paragraphs[0].style.name, "Title")
        self.assertTrue(any(p.style.name == "Heading 1" for p in document.paragraphs))
        self.assertTrue(any(p.style.name == "List Bullet" and long_argument.strip() in p.text for p in document.paragraphs))
        self.assertEqual(document.sections[0].page_width.inches, 8.5)
        self.assertEqual(document.sections[0].page_height.inches, 11)
        self.assertEqual(document.styles["Normal"].font.size.pt, 11)
        self.assertTrue(document.styles["Heading 1"].paragraph_format.keep_with_next)
        self.assertEqual(document.core_properties.author, "")

    def test_missing_dependency_reports_recoverable_export_error(self):
        original_import = __import__

        def unavailable(name, *args, **kwargs):
            if name == "docx" or name.startswith("docx."):
                raise ImportError("not installed")
            return original_import(name, *args, **kwargs)

        with patch("builtins.__import__", side_effect=unavailable):
            with self.assertRaisesRegex(InputError, "requirements.txt.*saved notes are unchanged"):
                docx_bytes(sample_draft())

    def test_title_border_removed_and_argument_label_kept_with_first_bullet(self):
        draft = sample_draft()
        original = copy.deepcopy(draft)
        document = Document(io.BytesIO(docx_bytes(draft)))
        self.assertIsNone(document.styles["Title"]._element.get_or_add_pPr().find(qn("w:pBdr")))
        labels = [paragraph for paragraph in document.paragraphs if paragraph.text == "Main Arguments:"]
        self.assertTrue(labels)
        self.assertTrue(all(paragraph.paragraph_format.keep_with_next for paragraph in labels))
        self.assertEqual(draft, original)

    def test_control_character_fails_clearly_without_changing_saved_content(self):
        draft = correct(sample_draft(), [{"path": "/testimonies/0/name", "value": "Test\x01Name"}])
        original = copy.deepcopy(draft)
        with self.assertRaisesRegex(InputError, "control character.*saved notes are unchanged"):
            docx_bytes(draft)
        self.assertEqual(draft, original)

    def test_shared_entry_keeps_each_testifier_and_confirmed_absence_of_questions(self):
        source = {
            "schema_version": 1, "synthetic": True,
            "hearing": {"id": "word-shared", "title": "Synthetic Shared Testimony", "date": "2026-10-05",
                        "committee": "Example Committee", "source_file": "synthetic-shared.mp3"},
            "segments": [
                {"id": "a1", "start": 0, "end": 5, "speaker": "Speaker A",
                 "text": "I'm Representative Vale. I am here to testify in support of H.1911."},
                {"id": "a2", "start": 6, "end": 10, "speaker": "Speaker A",
                 "text": "The bill requires express authority to create a trust."},
                {"id": "b1", "start": 11, "end": 15, "speaker": "Speaker B",
                 "text": "I'm Senator Ortega. I am here to testify in support of H.1911."},
                {"id": "b2", "start": 16, "end": 20, "speaker": "Speaker B",
                 "text": "It clarifies how families can plan for unforeseen circumstances."},
            ],
        }
        draft = correct(extract(source), operations=[{"op": "group_testimonies", "testimonies": [0, 1]}])
        text = content(docx_bytes(draft))
        self.assertEqual(text.count("Bill: H.1911"), 1)
        for expected in ("Representative Vale", "Senator Ortega", "Speaker A", "Speaker B",
                         "requires express authority", "unforeseen circumstances", "a2", "b2",
                         "Shared testimony [draft; review needed]", "individual speaker records retained"):
            self.assertIn(expected, text)
        self.assertNotIn("None noted during this testimony.", text)
        draft = correct(draft, operations=[
            {"op": "confirm_no_questions", "testimony": 0, "statement": 0, "confirmed": True, "evidence": ["a1", "a2"]},
            {"op": "confirm_no_questions", "testimony": 1, "statement": 0, "confirmed": True, "evidence": ["b1", "b2"]},
        ])
        text = content(docx_bytes(draft))
        self.assertEqual(text.count("None noted during this testimony. [human confirmed]"), 1)
        self.assertIn("00:00:00–00:00:05 · Speaker A · a1", text)
        self.assertIn("00:00:16–00:00:20 · Speaker B · b2", text)

    def test_shared_bill_keeps_later_testifiers_unapproved_state_visible(self):
        source = json.loads((ROOT / "fixtures/transcripts/multiple-bills.json").read_text(encoding="utf-8"))
        draft = correct(extract(source), [{"path": "/testimonies/1/statements/0/position", "value": "Support"}], operations=[
            {"op": "group_testimonies", "testimonies": [0, 1]},
            {"op": "confirm_no_questions", "testimony": 0, "statement": 1, "confirmed": True, "evidence": ["m05"]},
            {"op": "confirm_no_questions", "testimony": 1, "statement": 0, "confirmed": True, "evidence": ["m09"]},
        ])
        draft = correct(draft, [{"path": path, "review": "APPROVED"} for path, item in fields(draft)
                                if item["value"] and path != "/testimonies/1/statements/0/bill"
                                and not path.endswith("/questions_none")])
        text = content(docx_bytes(draft))
        self.assertIn("Review: Needs Review", text)
        self.assertIn("Bill: H.1001 [draft; review needed]", text)


if __name__ == "__main__":
    unittest.main()
